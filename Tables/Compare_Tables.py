"""
compare_table_structure.py

OCR-FREE table structure comparison.

Purpose:
    Compare English and translated PDF table structures.

Important:
    - Table location can change.
    - Table order is preserved.
    - Translation can cause a table to split across pages.
    - Repeated header rows are handled geometrically.
    - No OCR is used.

Example:

English:

    Page 10
        Table A = 2 columns, [8, 8]

Translated:

    Page 10
        Table A1 = 2 columns, [4, 4]

    Page 11
        Table A2 = 2 columns, [5, 5]

    A1 + A2 - repeated header
        = [8, 8]

Result:
    PASS
"""


import cv2
import numpy as np

from pathlib import Path
import json
import csv

import fitz  # PyMuPDF
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# PATH CONFIGURATION
# ============================================================

ENGLISH_DIR = Path(
    r"C:\Xylem Project\SpotCheck\Output\Table images\English"
)

TRANSLATED_DIR = Path(
     r"C:\Xylem Project\SpotCheck\Output\Table images\Translated-GE"
 )

# # Debug images showing detected tables
DEBUG_DIR = Path(
     r"C:\Xylem Project\SpotCheck\Output\Table images\Table comparison debug. English and GE"
 )

# Report files
REPORT_JSON = Path(
    r"C:\Xylem Project\SpotCheck\Output\Table images\Table comparison debug. English and GE\table_comparison.json"
)

REPORT_XLSX = Path(
    r"C:\Xylem Project\SpotCheck\Output\Table images\Table comparison debug. English and GE\table_comparison.xlsx"
)

# PDF Input files (for reading metadata page labels)
ENGLISH_PDF = Path(
    r"C:\Xylem Project\SpotCheck\Input\English\894387_5.0_en-US_2026-04_IOM.Start350.pdf"
)

TRANSLATED_PDF = Path(
   r"C:\Xylem Project\SpotCheck\Input\Translated\882549_5.0_el-GR_2026-04_IOM.Start350.pdf"
)


# ============================================================
# TABLE DETECTION SETTINGS
# ============================================================

# Minimum candidate table size in pixels
MIN_TABLE_WIDTH = 80
MIN_TABLE_HEIGHT = 25

# Minimum line length
MIN_HORIZONTAL_LINE_LENGTH = 20
MIN_VERTICAL_LINE_LENGTH = 15

# Morphological line detection
HORIZONTAL_KERNEL_DIVISOR = 40
VERTICAL_KERNEL_DIVISOR = 60

# How much of a column width must contain a horizontal
# line before we consider it a cell boundary.
HORIZONTAL_COVERAGE = 0.70

# How much of table height a vertical line should occupy
# to be considered a column boundary.
VERTICAL_COVERAGE = 0.25

# Merge nearby detected line positions
POSITION_TOLERANCE = 5

# Ignore very small gaps between boundaries
MIN_CELL_HEIGHT = 10


# ============================================================
# CONTINUATION SETTINGS
# ============================================================

# Previous fragment should finish reasonably close
# to the bottom of its page.
BOTTOM_PAGE_RATIO = 0.70

# Next fragment should start reasonably close
# to the top of the next page.
TOP_PAGE_RATIO = 0.30

# A fragment must have at least this many rows
# to be considered for continuation.
MIN_FRAGMENT_ROWS = 2

# Tolerance for row count differences.
#
# Translation can cause slight detection
# differences (e.g., text height changes).
#
# Allow ±1 row per column.
ROW_TOLERANCE = 1

# How many tables ahead to look
# when trying to re-sync after
# a mismatch.
RESYNC_LOOKAHEAD = 3


# ============================================================
# IMAGE FILES
# ============================================================

IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".tif",
    ".tiff"
}


# ============================================================
# LOAD IMAGE
# ============================================================

def load_image(path):

    image = cv2.imread(str(path))

    if image is None:
        raise RuntimeError(
            f"Could not read image:\n{path}"
        )

    return image


# ============================================================
# CREATE BINARY IMAGE
# ============================================================

def create_binary(image):

    gray = cv2.cvtColor(
        image,
        cv2.COLOR_BGR2GRAY
    )

    # Because your input is already masked:
    #
    # table lines/text = dark
    # background       = white
    #
    # Convert dark pixels to white.
    binary = cv2.threshold(
        gray,
        210,
        255,
        cv2.THRESH_BINARY_INV
    )[1]

    return binary


# ============================================================
# DETECT HORIZONTAL LINES
# ============================================================

def detect_horizontal_lines(binary):

    height, width = binary.shape

    kernel_length = max(
        MIN_HORIZONTAL_LINE_LENGTH,
        width // HORIZONTAL_KERNEL_DIVISOR
    )

    kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (
            kernel_length,
            1
        )
    )

    horizontal = cv2.morphologyEx(
        binary,
        cv2.MORPH_OPEN,
        kernel
    )

    return horizontal


# ============================================================
# DETECT VERTICAL LINES
# ============================================================

def detect_vertical_lines(binary):

    height, width = binary.shape

    kernel_length = max(
        MIN_VERTICAL_LINE_LENGTH,
        height // VERTICAL_KERNEL_DIVISOR
    )

    kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (
            1,
            kernel_length
        )
    )

    vertical = cv2.morphologyEx(
        binary,
        cv2.MORPH_OPEN,
        kernel
    )

    return vertical

# ============================================================
# CLUSTER NEARBY POSITIONS
# ============================================================

def cluster_positions(
    positions,
    tolerance=POSITION_TOLERANCE
):

    if not positions:
        return []

    positions = sorted(positions)

    clusters = [
        [positions[0]]
    ]

    for position in positions[1:]:

        if (
            position -
            clusters[-1][-1]
            <= tolerance
        ):
            clusters[-1].append(position)

        else:
            clusters.append(
                [position]
            )

    result = []

    for cluster in clusters:

        result.append(
            int(
                round(
                    sum(cluster)
                    /
                    len(cluster)
                )
            )
        )

    return result


# ============================================================
# DETECT TABLE REGIONS
# ============================================================

def detect_table_regions(
    horizontal,
    vertical
):

    combined = cv2.bitwise_or(
        horizontal,
        vertical
    )

    # Connect nearby line pieces.
    close_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (7, 7)
    )

    combined = cv2.morphologyEx(
        combined,
        cv2.MORPH_CLOSE,
        close_kernel
    )

    # Slightly expand the table structure.
    dilate_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (3, 3)
    )

    combined = cv2.dilate(
        combined,
        dilate_kernel,
        iterations=1
    )

    contours, _ = cv2.findContours(
        combined,
        cv2.RETR_EXTERNAL,
        cv2.CHAIN_APPROX_SIMPLE
    )

    page_h, page_w = combined.shape[:2]
    header_margin_px = int(page_h * 0.065)
    footer_margin_px = int(page_h * 0.065)

    regions = []

    for contour in contours:

        x, y, w, h = cv2.boundingRect(
            contour
        )

        if w < MIN_TABLE_WIDTH:
            continue

        if h < MIN_TABLE_HEIGHT:
            continue

        # Ignore candidates located in top header space
        if y < header_margin_px and (y + h) <= (header_margin_px + 30):
            continue

        # Ignore candidates located in bottom footer space
        if (y + h) > (page_h - footer_margin_px) and y >= (page_h - footer_margin_px - 30):
            continue

        regions.append(
            (
                x,
                y,
                w,
                h
            )
        )

    # Sort by vertical position.
    regions.sort(
        key=lambda r: (
            r[1],
            r[0]
        )
    )

    return regions


# ============================================================
# DETECT TABLE COLUMN BOUNDARIES
# ============================================================

def detect_vertical_boundaries(
    vertical,
    width,
    height
):

    projection = np.sum(
        vertical > 0,
        axis=0
    )

    required_length = max(
        MIN_VERTICAL_LINE_LENGTH,
        int(
            height *
            VERTICAL_COVERAGE
        )
    )

    positions = []

    for x, pixels in enumerate(
        projection
    ):

        if pixels >= required_length:

            positions.append(x)

    positions = cluster_positions(
        positions
    )

    if not positions:

        return [
            0,
            width - 1
        ]

    # Add left edge.
    if (
        positions[0]
        >
        POSITION_TOLERANCE
    ):

        positions.insert(
            0,
            0
        )

    else:

        positions[0] = 0

    # Add right edge.
    if (
        width - 1 -
        positions[-1]
        >
        POSITION_TOLERANCE
    ):

        positions.append(
            width - 1
        )

    else:

        positions[-1] = width - 1

    positions = clean_boundary_positions(
        positions,
        min_gap=25
    )

    return positions


# ============================================================
# DETECT HORIZONTAL BOUNDARIES
# FOR ONE COLUMN
# ============================================================

def detect_horizontal_boundaries_for_column(
    horizontal,
    x0,
    x1,
    height
):

    column = horizontal[
        :,
        x0:x1
    ]

    if column.shape[1] <= 0:

        return [
            0,
            height - 1
        ]

    projection = np.sum(
        column > 0,
        axis=1
    )

    column_width = x1 - x0

    required_length = max(
        MIN_HORIZONTAL_LINE_LENGTH,
        int(
            column_width *
            HORIZONTAL_COVERAGE
        )
    )

    positions = []

    for y, pixels in enumerate(
        projection
    ):

        if pixels >= required_length:

            positions.append(y)

    positions = cluster_positions(
        positions
    )

    if not positions:

        return [
            0,
            height - 1
        ]

    # Top boundary.
    if (
        positions[0]
        >
        POSITION_TOLERANCE
    ):

        positions.insert(
            0,
            0
        )

    else:

        positions[0] = 0

    # Bottom boundary.
    if (
        height - 1 -
        positions[-1]
        >
        POSITION_TOLERANCE
    ):

        positions.append(
            height - 1
        )

    else:

        positions[-1] = height - 1

    return positions


# ============================================================
# CLEAN SMALL GAPS
# ============================================================

def clean_boundary_positions(
    positions,
    min_gap=MIN_CELL_HEIGHT
):

    if len(positions) <= 2:

        return positions

    cleaned = [
        positions[0]
    ]

    for position in positions[1:-1]:

        previous = cleaned[-1]

        if (
            position -
            previous
            >= min_gap
        ):

            cleaned.append(
                position
            )

    # Always preserve final boundary.
    if (
        positions[-1] -
        cleaned[-1]
        >= min_gap
    ):

        cleaned.append(
            positions[-1]
        )

    return cleaned


# ============================================================
# DETECT TEXT OVERFLOW FROM TABLE
# ============================================================

def detect_text_overflow(
    binary,
    horizontal,
    vertical,
    region,
    vertical_boundaries,
    horizontal_boundaries
):

    if binary is None:
        return {
            "has_overflow": False,
            "overflow_reasons": [],
            "overflow_bboxes": []
        }

    x, y, w, h = region
    img_height, img_width = binary.shape

    # Isolate text pixels (remove horizontal and vertical lines)
    lines_mask = cv2.bitwise_or(horizontal, vertical)
    text_mask = cv2.bitwise_and(binary, cv2.bitwise_not(lines_mask))

    overflow_found = False
    reasons = []
    overflow_bboxes = []

    # 1. Outer Bottom Overflow (Text spilling below table bottom border)
    margin_bottom = 20
    y_start = min(img_height, y + h)
    y_end = min(img_height, y + h + margin_bottom)
    x_start = max(0, x + 5)
    x_end = min(img_width, x + w - 5)

    if y_end > y_start and x_end > x_start:
        bottom_region = text_mask[y_start:y_end, x_start:x_end]
        contours, _ = cv2.findContours(
            bottom_region, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        for c in contours:
            area = cv2.contourArea(c)
            if area >= 12:  # Min text character area
                bx, by, bw, bh = cv2.boundingRect(c)
                overflow_found = True
                reasons.append("Text spills below bottom table border")
                overflow_bboxes.append(
                    (x_start + bx, y_start + by, bw, bh)
                )
                break

    # 2. Outer Right Overflow (Text spilling past right border)
    margin_side = 20
    rx_start = min(img_width, x + w)
    rx_end = min(img_width, x + w + margin_side)
    ry_start = max(0, y + 5)
    ry_end = min(img_height, y + h - 5)

    if rx_end > rx_start and ry_end > ry_start:
        right_region = text_mask[ry_start:ry_end, rx_start:rx_end]
        contours, _ = cv2.findContours(
            right_region, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        for c in contours:
            area = cv2.contourArea(c)
            if area >= 12:
                bx, by, bw, bh = cv2.boundingRect(c)
                overflow_found = True
                reasons.append("Text spills past right table border")
                overflow_bboxes.append(
                    (rx_start + bx, ry_start + by, bw, bh)
                )
                break

    # 3. Outer Left Overflow (Text spilling past left border)
    lx_start = max(0, x - margin_side)
    lx_end = max(0, x)

    if lx_end > lx_start and ry_end > ry_start:
        left_region = text_mask[ry_start:ry_end, lx_start:lx_end]
        contours, _ = cv2.findContours(
            left_region, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        for c in contours:
            area = cv2.contourArea(c)
            if area >= 12:
                bx, by, bw, bh = cv2.boundingRect(c)
                overflow_found = True
                reasons.append("Text spills past left table border")
                overflow_bboxes.append(
                    (lx_start + bx, ry_start + by, bw, bh)
                )
                break

    # 4. Cell Line Collision Overflow (Text colliding directly with bottom horizontal line of cell)
    t_crop_text = text_mask[y:y+h, x:x+w]
    for col_idx, boundaries in enumerate(horizontal_boundaries):
        if col_idx >= len(vertical_boundaries) - 1:
            continue
        cx0 = vertical_boundaries[col_idx]
        cx1 = vertical_boundaries[col_idx + 1]

        for b_idx in range(len(boundaries) - 1):
            by1 = boundaries[b_idx + 1]
            if by1 <= 3 or by1 >= h - 1:
                continue

            strip_y0 = max(0, by1 - 3)
            strip_y1 = by1
            strip = t_crop_text[strip_y0:strip_y1, cx0+2:cx1-2]

            if strip.shape[0] > 0 and strip.shape[1] > 0:
                text_pixels = np.sum(strip > 0)
                if text_pixels >= 15:
                    overflow_found = True
                    reasons.append(f"Text collides with cell bottom boundary (Col {col_idx+1}, Row {b_idx+1})")
                    overflow_bboxes.append(
                        (x + cx0 + 2, y + strip_y0, cx1 - cx0 - 4, 4)
                    )
                    break

    unique_reasons = list(set(reasons))

    return {
        "has_overflow": overflow_found,
        "overflow_reasons": unique_reasons,
        "overflow_bboxes": overflow_bboxes
    }


# ============================================================
# ANALYZE ONE TABLE
# ============================================================

def analyze_table(
    horizontal,
    vertical,
    region,
    binary=None
):

    x, y, w, h = region

    h_crop = horizontal[
        y:y + h,
        x:x + w
    ]

    v_crop = vertical[
        y:y + h,
        x:x + w
    ]

    # --------------------------------------------------------
    # COLUMN BOUNDARIES
    # --------------------------------------------------------

    vertical_boundaries = (
        detect_vertical_boundaries(
            v_crop,
            w,
            h
        )
    )

    if len(
        vertical_boundaries
    ) < 2:

        return None

    columns = (
        len(vertical_boundaries)
        - 1
    )

    # --------------------------------------------------------
    # ROWS PER COLUMN
    # --------------------------------------------------------

    rows_per_column = []

    horizontal_boundaries = []

    for column_index in range(
        columns
    ):

        x0 = vertical_boundaries[
            column_index
        ]

        x1 = vertical_boundaries[
            column_index + 1
        ]

        boundaries = (
            detect_horizontal_boundaries_for_column(
                h_crop,
                x0,
                x1,
                h
            )
        )

        boundaries = (
            clean_boundary_positions(
                boundaries
            )
        )

        row_count = max(
            1,
            len(boundaries) - 1
        )

        rows_per_column.append(
            row_count
        )

        horizontal_boundaries.append(
            boundaries
        )

    # Detect text overflow
    overflow_info = detect_text_overflow(
        binary,
        horizontal,
        vertical,
        region,
        vertical_boundaries,
        horizontal_boundaries
    )

    return {
        "bbox": region,

        "columns": columns,

        "rows_per_column":
            rows_per_column,

        "vertical_boundaries":
            vertical_boundaries,

        "horizontal_boundaries":
            horizontal_boundaries,

        "has_text_overflow":
            overflow_info["has_overflow"],

        "overflow_reasons":
            overflow_info["overflow_reasons"],

        "overflow_bboxes":
            overflow_info["overflow_bboxes"]
    }


# ============================================================
# PAGE POSITION HELPERS
# ============================================================

def get_page_position(
    table
):

    page_height = table["page_height"]

    x, y, w, h = table["bbox"]

    top_ratio = (
        y /
        page_height
    )

    bottom_ratio = (
        (y + h) /
        page_height
    )

    return (
        top_ratio,
        bottom_ratio
    )


def is_bottom_table(
    table
):

    _, bottom_ratio = (
        get_page_position(
            table
        )
    )

    return (
        bottom_ratio
        >=
        BOTTOM_PAGE_RATIO
    )


def is_top_table(
    table
):

    top_ratio, _ = (
        get_page_position(
            table
        )
    )

    return (
        top_ratio
        <=
        TOP_PAGE_RATIO
    )


# ============================================================
# TABLE SIGNATURE
# ============================================================

def table_signature(
    table
):

    return {
        "columns":
            table["columns"],

        "rows_per_column":
            list(
                table[
                    "rows_per_column"
                ]
            )
    }


# ============================================================
# EXACT TABLE MATCH
# ============================================================

def tables_match(
    table1,
    table2
):

    return (
        table1["columns"]
        ==
        table2["columns"]
        and
        table1[
            "rows_per_column"
        ]
        ==
        table2[
            "rows_per_column"
        ]
    )


# ============================================================
# TOLERANT TABLE MATCH
#
# Same column count required.
# Rows per column may differ by
# ± ROW_TOLERANCE.
# ============================================================

def tables_match_tolerant(
    table1,
    table2
):

    if (
        table1["columns"]
        !=
        table2["columns"]
    ):

        return False

    rows1 = table1[
        "rows_per_column"
    ]

    rows2 = table2[
        "rows_per_column"
    ]

    if len(rows1) != len(rows2):
        return False

    for r1, r2 in zip(
        rows1,
        rows2
    ):

        if (
            abs(r1 - r2)
            >
            ROW_TOLERANCE
        ):

            return False

    return True


# ============================================================
# MERGE TWO FRAGMENTS
#
# repeated_header = False
#     rows1 + rows2
#
# repeated_header = True
#     rows1 + rows2 - 1
#
# ============================================================

def merge_fragments(
    table1,
    table2,
    repeated_header
):

    if (
        table1["columns"]
        !=
        table2["columns"]
    ):

        return None

    rows1 = table1[
        "rows_per_column"
    ]

    rows2 = table2[
        "rows_per_column"
    ]

    merged_rows = []

    for r1, r2 in zip(
        rows1,
        rows2
    ):

        if repeated_header:

            value = (
                r1 +
                r2 -
                1
            )

        else:

            value = (
                r1 +
                r2
            )

        merged_rows.append(
            value
        )

    merged = {
        "columns":
            table1["columns"],

        "rows_per_column":
            merged_rows
    }

    return merged


# ============================================================
# CAN TWO TABLES BE CONTINUATION?
# ============================================================

def is_last_table_on_page(
    table,
    all_tables
):

    page = table["page_index"]

    same_page = [
        t for t in all_tables
        if t["page_index"] == page
    ]

    if not same_page:
        return False

    last = max(
        same_page,
        key=lambda t: (
            t["physical_table_index"]
        )
    )

    return (
        table["physical_table_index"]
        ==
        last["physical_table_index"]
    )


def is_first_table_on_page(
    table,
    all_tables
):

    page = table["page_index"]

    same_page = [
        t for t in all_tables
        if t["page_index"] == page
    ]

    if not same_page:
        return False

    first = min(
        same_page,
        key=lambda t: (
            t["physical_table_index"]
        )
    )

    return (
        table["physical_table_index"]
        ==
        first["physical_table_index"]
    )


def possible_continuation(
    table1,
    table2,
    all_tables
):

    # --------------------------------------------------------
    # Must be consecutive pages.
    # --------------------------------------------------------

    if (
        table2["page_index"]
        !=
        table1["page_index"] + 1
    ):

        return False

    # --------------------------------------------------------
    # Same number of columns.
    # --------------------------------------------------------

    if (
        table1["columns"]
        !=
        table2["columns"]
    ):

        return False

    # --------------------------------------------------------
    # Previous table near bottom.
    # --------------------------------------------------------

    if not is_bottom_table(
        table1
    ):

        return False

    # --------------------------------------------------------
    # Next table near top.
    # --------------------------------------------------------

    if not is_top_table(
        table2
    ):

        return False

    # --------------------------------------------------------
    # First fragment must be the last table
    # on its page.
    # --------------------------------------------------------

    if not is_last_table_on_page(
        table1,
        all_tables
    ):

        return False

    # --------------------------------------------------------
    # Second fragment must be the first table
    # on its page.
    # --------------------------------------------------------

    if not is_first_table_on_page(
        table2,
        all_tables
    ):

        return False

    # --------------------------------------------------------
    # Both should have meaningful number of rows.
    # --------------------------------------------------------

    rows1 = table1[
        "rows_per_column"
    ]

    rows2 = table2[
        "rows_per_column"
    ]

    if max(rows1) < MIN_FRAGMENT_ROWS:
        return False

    if max(rows2) < MIN_FRAGMENT_ROWS:
        return False

    return True


# ============================================================
# TRY TO MATCH A SPLIT TABLE
# ============================================================

def find_split_match(
    target_table,
    fragment1,
    fragment2
):

    # Same number of columns is mandatory.
    if (
        target_table["columns"]
        !=
        fragment1["columns"]
        or
        target_table["columns"]
        !=
        fragment2["columns"]
    ):

        return None

    # --------------------------------------------------------
    # OPTION 1
    #
    # No repeated header.
    #
    # rows1 + rows2
    # --------------------------------------------------------

    merged_no_header = (
        merge_fragments(
            fragment1,
            fragment2,
            repeated_header=False
        )
    )

    if (
        merged_no_header
        and
        merged_no_header[
            "rows_per_column"
        ]
        ==
        target_table[
            "rows_per_column"
        ]
    ):

        return {
            "matched": True,
            "repeated_header": False,
            "merged":
                merged_no_header
        }

    # --------------------------------------------------------
    # OPTION 2
    #
    # Repeated header.
    #
    # rows1 + rows2 - 1
    # --------------------------------------------------------

    merged_with_header = (
        merge_fragments(
            fragment1,
            fragment2,
            repeated_header=True
        )
    )

    if (
        merged_with_header
        and
        merged_with_header[
            "rows_per_column"
        ]
        ==
        target_table[
            "rows_per_column"
        ]
    ):

        return {
            "matched": True,
            "repeated_header": True,
            "merged":
                merged_with_header
        }

    return None


# ============================================================
# MERGE MULTIPLE FRAGMENTS
#
# Merges a list of consecutive fragments.
#
# repeated_header = True
#     Subtracts 1 row per continuation fragment
#     (all except the first).
#
# repeated_header = False
#     Simple sum of all rows.
# ============================================================

def merge_multiple_fragments(
    fragments,
    repeated_header
):

    if not fragments:
        return None

    if len(fragments) == 1:

        return {
            "columns":
                fragments[0]["columns"],

            "rows_per_column":
                list(
                    fragments[0][
                        "rows_per_column"
                    ]
                )
        }

    # All fragments must have
    # the same column count.
    columns = fragments[0]["columns"]

    for fragment in fragments[1:]:

        if (
            fragment["columns"]
            !=
            columns
        ):

            return None

    merged_rows = list(
        fragments[0][
            "rows_per_column"
        ]
    )

    for fragment in fragments[1:]:

        rows = fragment[
            "rows_per_column"
        ]

        for i in range(
            len(merged_rows)
        ):

            if repeated_header:

                merged_rows[i] += (
                    rows[i] - 1
                )

            else:

                merged_rows[i] += (
                    rows[i]
                )

    return {
        "columns": columns,

        "rows_per_column":
            merged_rows
    }


# ============================================================
# FIND MANY-TO-MANY MULTI-FRAGMENT SPLIT MATCH
#
# Handles cases where tables are split across pages on
# BOTH English and Translated sides (e.g. 2 pages vs 3 pages).
# ============================================================

def find_many_to_many_split_match(
    english_tables,
    eng_start,
    translated_tables,
    trans_start
):

    # Gather consecutive continuation fragments for English
    eng_chain = [
        english_tables[eng_start]
    ]

    idx = eng_start + 1

    while idx < len(english_tables):

        if possible_continuation(
            eng_chain[-1],
            english_tables[idx],
            english_tables
        ):

            eng_chain.append(
                english_tables[idx]
            )

            idx += 1

        else:

            break

    # Gather consecutive continuation fragments for Translated
    trans_chain = [
        translated_tables[trans_start]
    ]

    idx = trans_start + 1

    while idx < len(translated_tables):

        if possible_continuation(
            trans_chain[-1],
            translated_tables[idx],
            translated_tables
        ):

            trans_chain.append(
                translated_tables[idx]
            )

            idx += 1

        else:

            break

    # If neither side has continuation fragments, no split match possible
    if len(eng_chain) <= 1 and len(trans_chain) <= 1:
        return None

    # Search combinations of (eng_count, trans_count)
    # Prefer larger counts first to capture full split tables
    for eng_count in range(len(eng_chain), 0, -1):

        for trans_count in range(len(trans_chain), 0, -1):

            if eng_count == 1 and trans_count == 1:
                continue

            eng_sub = eng_chain[:eng_count]
            trans_sub = trans_chain[:trans_count]

            for eng_header in [True, False]:

                merged_eng = merge_multiple_fragments(
                    eng_sub,
                    repeated_header=(eng_header if eng_count > 1 else False)
                )

                if not merged_eng:
                    continue

                for trans_header in [True, False]:

                    merged_trans = merge_multiple_fragments(
                        trans_sub,
                        repeated_header=(trans_header if trans_count > 1 else False)
                    )

                    if not merged_trans:
                        continue

                    if (
                        tables_match(
                            merged_eng,
                            merged_trans
                        )
                        or
                        tables_match_tolerant(
                            merged_eng,
                            merged_trans
                        )
                    ):

                        return {
                            "matched": True,
                            "eng_fragments": eng_sub,
                            "trans_fragments": trans_sub,
                            "eng_count": eng_count,
                            "trans_count": trans_count,
                            "merged": merged_eng,
                            "repeated_header": (
                                eng_header if eng_count > 1 else trans_header
                            )
                        }

    return None


# ============================================================
# ANALYZE ONE PAGE
# ============================================================

def analyze_page(
    image_path,
    page_index,
    debug=False
):

    image = load_image(
        image_path
    )

    binary = create_binary(
        image
    )

    horizontal = (
        detect_horizontal_lines(
            binary
        )
    )

    vertical = (
        detect_vertical_lines(
            binary
        )
    )

    regions = detect_table_regions(
        horizontal,
        vertical
    )

    tables = []

    debug_image = image.copy()

    for table_index, region in enumerate(
        regions,
        start=1
    ):

        table = analyze_table(
            horizontal,
            vertical,
            region,
            binary=binary
        )

        if table is None:
            continue

        table["page_index"] = page_index
        table["image_path"] = str(
            image_path
        )

        table["page_height"] = (
            image.shape[0]
        )

        table["physical_table_index"] = (
            table_index
        )

        tables.append(
            table
        )

        # ----------------------------------------------------
        # DEBUG IMAGE
        # ----------------------------------------------------

        if debug:

            x, y, w, h = region

            border_color = (0, 0, 255)

            if table.get("has_text_overflow"):
                border_color = (0, 140, 255)  # Orange/Amber for overflow

                # Highlight table boundary for overflow
                cv2.rectangle(
                    debug_image,
                    (x - 2, y - 2),
                    (x + w + 2, y + h + 2),
                    (0, 140, 255),
                    3
                )

                # Draw specific overflow boxes
                for obx, oby, obw, obh in table.get("overflow_bboxes", []):
                    cv2.rectangle(
                        debug_image,
                        (obx, oby),
                        (obx + obw, oby + obh),
                        (0, 0, 255),
                        2
                    )

            cv2.rectangle(
                debug_image,
                (x, y),
                (x + w, y + h),
                border_color,
                2
            )

            # Column boundaries
            for bx in table[
                "vertical_boundaries"
            ]:

                cv2.line(
                    debug_image,
                    (x + bx, y),
                    (x + bx, y + h),
                    (255, 0, 0),
                    1
                )

            # Horizontal boundaries
            for col, boundaries in enumerate(
                table[
                    "horizontal_boundaries"
                ]
            ):

                cx0 = table[
                    "vertical_boundaries"
                ][col]

                cx1 = table[
                    "vertical_boundaries"
                ][col + 1]

                for by in boundaries:

                    cv2.line(
                        debug_image,
                        (
                            x + cx0,
                            y + by
                        ),
                        (
                            x + cx1,
                            y + by
                        ),
                        (0, 255, 0),
                        1
                    )

            label = (
                f"T{table_index}: "
                f"{table['columns']}C "
                f"{table['rows_per_column']}"
            )

            if table.get("has_text_overflow"):
                label += " [TEXT OVERFLOW!]"

            cv2.putText(
                debug_image,
                label,
                (
                    x,
                    max(
                        20,
                        y - 5
                    )
                ),
                cv2.FONT_HERSHEY_SIMPLEX,
                0.55,
                (0, 140, 255) if table.get("has_text_overflow") else (0, 0, 255),
                2
            )

    if debug:

        DEBUG_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        debug_path = (
            DEBUG_DIR /
            image_path.name
        )

        cv2.imwrite(
            str(debug_path),
            debug_image
        )

    return tables


# ============================================================
# ANALYZE COMPLETE FOLDER
# ============================================================

def analyze_folder(
    folder
):

    if not folder.exists():

        raise FileNotFoundError(
            f"Folder does not exist:\n{folder}"
        )

    image_paths = sorted(
        [
            p
            for p in folder.iterdir()
            if p.suffix.lower()
            in IMAGE_EXTENSIONS
        ]
    )

    all_tables = []

    for page_index, image_path in enumerate(
        image_paths
    ):

        print(
            f"Analyzing page "
            f"{page_index + 1}: "
            f"{image_path.name}"
        )

        tables = analyze_page(
            image_path,
            page_index,
            debug=True
        )

        all_tables.extend(
            tables
        )

    return all_tables


# ============================================================
# COMPARE ENGLISH -> TRANSLATED
#
# The English sequence is the reference.
#
# Supports multi-fragment splits (3+ pages)
# in both directions.
# ============================================================

def compare_direction(
    english_tables,
    translated_tables
):

    results = []

    english_index = 0
    translated_index = 0

    while (
        english_index
        <
        len(english_tables)
    ):

        english = english_tables[
            english_index
        ]

        # ----------------------------------------------------
        # No translated table left.
        # ----------------------------------------------------

        if (
            translated_index
            >=
            len(translated_tables)
        ):

            results.append({
                "table_number":
                    english_index + 1,

                "status":
                    "MISSING",

                "english":
                    english,

                "translated":
                    None
            })

            english_index += 1

            continue

        translated = translated_tables[
            translated_index
        ]

        # ====================================================
        # 1. CROSS-PAGE SPLIT MATCH (N-to-M Fragments)
        #
        # Table split across pages on English, Translated,
        # or both (checked first to capture full multi-page
        # continuation chains).
        # ====================================================

        split_match = (
            find_many_to_many_split_match(
                english_tables,
                english_index,
                translated_tables,
                translated_index
            )
        )

        if split_match is not None:

            eng_count = split_match["eng_count"]
            trans_count = split_match["trans_count"]

            results.append({
                "table_number":
                    english_index + 1,

                "status":
                    "PASS",

                "reason":
                    f"Cross-page table split match "
                    f"(English {eng_count} page(s), "
                    f"Translated {trans_count} page(s))",

                "english":
                    (
                        english
                        if eng_count == 1
                        else None
                    ),

                "english_fragments":
                    split_match["eng_fragments"],

                "translated_fragments":
                    split_match["trans_fragments"],

                "merged":
                    split_match["merged"],

                "repeated_header":
                    split_match["repeated_header"]
            })

            english_index += eng_count
            translated_index += trans_count

            continue

        # ====================================================
        # 2. EXACT 1-TO-1 MATCH
        # ====================================================

        if tables_match(
            english,
            translated
        ):

            results.append({
                "table_number":
                    english_index + 1,

                "status":
                    "PASS",

                "reason":
                    "Exact structure match",

                "english":
                    english,

                "translated_fragments":
                    [translated]
            })

            english_index += 1
            translated_index += 1

            continue

        # ====================================================
        # 4. TOLERANT MATCH
        #
        # Same column count, row counts
        # within ± ROW_TOLERANCE.
        # ====================================================

        if tables_match_tolerant(
            english,
            translated
        ):

            results.append({
                "table_number":
                    english_index + 1,

                "status":
                    "PASS",

                "reason":
                    "Structure match "
                    "within tolerance",

                "english":
                    english,

                "translated_fragments":
                    [translated]
            })

            english_index += 1
            translated_index += 1

            continue

        # ====================================================
        # 5. RE-SYNC
        #
        # A mismatch occurred.
        #
        # Before marking DIFFERENT and
        # advancing both, try:
        #
        #   A) Skip translated table
        #      (extra/phantom table)
        #
        #   B) Skip English table
        #      (missing translated table)
        #
        # This prevents a single detection
        # error from cascading through
        # all subsequent comparisons.
        # ====================================================

        resynced = False

        # ------------------------------------------------
        # A) Try skipping translated tables
        #    to find one that matches English.
        # ------------------------------------------------

        for skip in range(
            1,
            RESYNC_LOOKAHEAD + 1
        ):

            ahead_index = (
                translated_index + skip
            )

            if (
                ahead_index
                >=
                len(translated_tables)
            ):

                break

            ahead = translated_tables[
                ahead_index
            ]

            if (
                tables_match(
                    english,
                    ahead
                )
                or
                tables_match_tolerant(
                    english,
                    ahead
                )
            ):

                # Mark skipped translated
                # tables as EXTRA.
                for extra_i in range(
                    translated_index,
                    ahead_index
                ):

                    results.append({
                        "table_number":
                            len(results)
                            + 1,

                        "status":
                            "EXTRA",

                        "reason":
                            "Extra translated "
                            "table (re-sync)",

                        "english":
                            None,

                        "translated_fragments":
                            [
                                translated_tables[
                                    extra_i
                                ]
                            ]
                    })

                translated_index = (
                    ahead_index
                )

                resynced = True

                break

        if resynced:
            continue

        # ------------------------------------------------
        # B) Try skipping English tables
        #    to find one that matches translated.
        # ------------------------------------------------

        for skip in range(
            1,
            RESYNC_LOOKAHEAD + 1
        ):

            ahead_index = (
                english_index + skip
            )

            if (
                ahead_index
                >=
                len(english_tables)
            ):

                break

            ahead = english_tables[
                ahead_index
            ]

            if (
                tables_match(
                    ahead,
                    translated
                )
                or
                tables_match_tolerant(
                    ahead,
                    translated
                )
            ):

                # Mark skipped English
                # tables as MISSING.
                for miss_i in range(
                    english_index,
                    ahead_index
                ):

                    results.append({
                        "table_number":
                            miss_i + 1,

                        "status":
                            "MISSING",

                        "reason":
                            "No matching "
                            "translated table "
                            "(re-sync)",

                        "english":
                            english_tables[
                                miss_i
                            ],

                        "translated":
                            None
                    })

                english_index = (
                    ahead_index
                )

                resynced = True

                break

        if resynced:
            continue

        # ====================================================
        # 6. DIFFERENT
        #
        # No re-sync possible.
        # ====================================================

        results.append({
            "table_number":
                english_index + 1,

            "status":
                "DIFFERENT",

            "reason":
                "Column/row structure "
                "does not match",

            "english":
                english,

            "translated_fragments":
                [translated]
        })

        english_index += 1
        translated_index += 1

    # ========================================================
    # EXTRA TRANSLATED TABLES
    # ========================================================

    while (
        translated_index
        <
        len(translated_tables)
    ):

        results.append({
            "table_number":
                len(results) + 1,

            "status":
                "EXTRA",

            "reason":
                "Extra translated table",

            "english":
                None,

            "translated_fragments":
                [
                    translated_tables[
                        translated_index
                    ]
                ]
        })

        translated_index += 1

    return results


# ============================================================
# FORMAT TABLE
# ============================================================

def format_table(
    table
):

    if table is None:

        return None

    res = {
        "page_index":
            table.get(
                "page_index"
            ),

        "page":
            Path(
                table.get(
                    "image_path",
                    ""
                )
            ).name,

        "table_index":
            table.get(
                "physical_table_index"
            ),

        "columns":
            table[
                "columns"
            ],

        "rows_per_column":
            table[
                "rows_per_column"
            ],

        "bbox":
            table[
                "bbox"
            ]
    }

    if table.get("has_text_overflow"):

        res["has_text_overflow"] = True
        res["overflow_reasons"] = table.get("overflow_reasons", [])

    return res


# ============================================================
# LOAD PDF PAGE LABELS
# ============================================================

def load_pdf_page_labels(
    pdf_path
):

    labels = {}

    if not pdf_path or not Path(pdf_path).exists():
        return labels

    try:
        doc = fitz.open(pdf_path)

        for i in range(len(doc)):

            lbl = doc[i].get_label()

            labels[i] = str(lbl) if lbl else str(i + 1)

    except Exception as e:

        print(
            f"Warning: Could not read PDF page labels from {pdf_path}: {e}"
        )

    return labels


# ============================================================
# CONVERT RESULT TO JSON-SAFE & ENRICHED REPORT
# ============================================================

def create_report(
    results,
    english_labels=None,
    translated_labels=None
):

    if english_labels is None:
        english_labels = {}

    if translated_labels is None:
        translated_labels = {}

    report = []

    for result in results:

        row = {
            "table_number":
                result[
                    "table_number"
                ],

            "status":
                result[
                    "status"
                ],

            "reason":
                result.get(
                    "reason",
                    ""
                )
        }

        # Determine English and Translated Fragments
        eng_frags = []
        if result.get("english"):
            eng_frags.append(format_table(result["english"]))
        elif result.get("english_fragments"):
            eng_frags.extend([format_table(t) for t in result["english_fragments"]])

        trans_frags = []
        if result.get("translated_fragments"):
            trans_frags.extend([format_table(t) for t in result["translated_fragments"]])

        if "english" in result and result["english"]:

            row["english"] = format_table(
                result["english"]
            )

        if "english_fragments" in result:

            row[
                "english_fragments"
            ] = [
                format_table(t)
                for t in result[
                    "english_fragments"
                ]
            ]

        if "translated_fragments" in result:

            row[
                "translated_fragments"
            ] = [
                format_table(t)
                for t in result[
                    "translated_fragments"
                ]
            ]

        if "merged" in result:

            row["merged"] = result[
                "merged"
            ]

        row[
            "repeated_header"
        ] = result.get(
            "repeated_header",
            False
        )

        # ----------------------------------------------------
        # SEQUENCE PAGE NUMBER (Physical start & end pages)
        # ----------------------------------------------------
        eng_seq = ""
        if eng_frags:
            es_idx = eng_frags[0]["page_index"] + 1
            ee_idx = eng_frags[-1]["page_index"] + 1
            eng_seq = f"Eng P{es_idx}" if es_idx == ee_idx else f"Eng P{es_idx}-P{ee_idx}"

        trans_seq = ""
        if trans_frags:
            ts_idx = trans_frags[0]["page_index"] + 1
            te_idx = trans_frags[-1]["page_index"] + 1
            trans_seq = f"Trans P{ts_idx}" if ts_idx == te_idx else f"Trans P{ts_idx}-P{te_idx}"

        if eng_seq and trans_seq:
            sequence_page_number = f"{eng_seq} | {trans_seq}"
        else:
            sequence_page_number = eng_seq or trans_seq

        row["sequence_page_number"] = sequence_page_number

        # ----------------------------------------------------
        # LABELED PAGE NUMBER (Printed page labels from PDF)
        # ----------------------------------------------------
        eng_lbl = ""
        if eng_frags:
            es_pi = eng_frags[0]["page_index"]
            ee_pi = eng_frags[-1]["page_index"]
            el_start = english_labels.get(es_pi, str(es_pi + 1))
            el_end = english_labels.get(ee_pi, str(ee_pi + 1))
            eng_lbl = f"Eng {el_start}" if el_start == el_end else f"Eng {el_start}-{el_end}"

        trans_lbl = ""
        if trans_frags:
            ts_pi = trans_frags[0]["page_index"]
            te_pi = trans_frags[-1]["page_index"]
            tl_start = translated_labels.get(ts_pi, str(ts_pi + 1))
            tl_end = translated_labels.get(te_pi, str(te_pi + 1))
            trans_lbl = f"Trans {tl_start}" if tl_start == tl_end else f"Trans {tl_start}-{tl_end}"

        if eng_lbl and trans_lbl:
            labeled_page_number = f"{eng_lbl} | {trans_lbl}"
        else:
            labeled_page_number = eng_lbl or trans_lbl

        row["labeled_page_number"] = labeled_page_number

        # Gather overflow status from english and translated fragments
        all_tables_in_result = []
        if result.get("english"):
            all_tables_in_result.append(result["english"])
        if result.get("english_fragments"):
            all_tables_in_result.extend(result["english_fragments"])
        if result.get("translated_fragments"):
            all_tables_in_result.extend(result["translated_fragments"])

        has_overflow = False
        all_overflow_reasons = []

        for t in all_tables_in_result:
            if t and t.get("has_text_overflow"):
                has_overflow = True
                all_overflow_reasons.extend(t.get("overflow_reasons", []))

        row["has_text_overflow"] = has_overflow
        row["overflow_reasons"] = list(set(all_overflow_reasons))
        row["overflow_reasons_str"] = "; ".join(row["overflow_reasons"])

        report.append(row)

    return report


# ============================================================
# SAVE JSON
# ============================================================

def save_json(
    report
):

    REPORT_JSON.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    try:

        with open(
            REPORT_JSON,
            "w",
            encoding="utf-8"
        ) as f:

            json.dump(
                report,
                f,
                indent=4
            )

    except PermissionError:

        print(
            f"Warning: Could not write JSON report (file locked): {REPORT_JSON}"
        )


# ============================================================
# SAVE EXCEL (STYLISH WITH ROW COLORING)
# ============================================================

def save_excel(
    report
):

    REPORT_XLSX.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table Comparison"

    headers = [
        "Table Number",
        "Status",
        "Reason",
        "Sequence Page Number",
        "Labeled Page Number",
        "English Columns",
        "English Rows",
        "Translated Columns",
        "Translated Rows",
        "Repeated Header",
        "Text Overflow",
        "Overflow Details"
    ]

    ws.append(headers)

    # --------------------------------------------------------
    # STYLES & COLORS
    # --------------------------------------------------------
    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    header_font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color="FFFFFF"
    )

    align_center = Alignment(
        horizontal="center",
        vertical="center"
    )

    align_left = Alignment(
        horizontal="left",
        vertical="center"
    )

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Header Row Styling
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Row Fills:
    # 1. Green (PASS + No Overflow)
    green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    green_font = Font(name="Calibri", size=10, color="155724", bold=True)

    # 2. Yellow/Amber (PASS + Text Overflow Warning -> Review Required)
    amber_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    amber_font = Font(name="Calibri", size=10, color="856404", bold=True)

    # 3. Red/Pink (DIFFERENT/MISSING/EXTRA -> Action Required)
    red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    red_font = Font(name="Calibri", size=10, color="721C24", bold=True)

    normal_font = Font(name="Calibri", size=10, color="000000")

    for row_idx, item in enumerate(report, start=2):

        status = item.get("status", "")
        overflow = item.get("has_text_overflow", False)

        if overflow:
            row_fill = amber_fill
            status_font = amber_font
            display_status = f"{status} (Review Overflow)"
        elif status == "PASS":
            row_fill = green_fill
            status_font = green_font
            display_status = "PASS (OK)"
        else:
            row_fill = red_fill
            status_font = red_font
            display_status = f"{status} (Review Structure)"

        english_frags = item.get("english_fragments", [])
        english = item.get("english")
        translated_frags = item.get("translated_fragments", [])
        merged = item.get("merged")

        # English columns / rows description
        if english_frags and len(english_frags) > 1:
            eng_cols = f"{english_frags[0]['columns']} ({len(english_frags)} fragments)"
            eng_detail = " + ".join(f"P{t['page_index']+1}{t['rows_per_column']}" for t in english_frags)
            eng_rows = f"{merged['rows_per_column'] if merged else ''} (Combined: {eng_detail})"
        elif english:
            eng_cols = english["columns"]
            eng_rows = str(english["rows_per_column"])
        elif english_frags:
            eng_cols = english_frags[0]["columns"]
            eng_rows = str(english_frags[0]["rows_per_column"])
        else:
            eng_cols = ""
            eng_rows = ""

        # Translated columns / rows description
        if translated_frags and len(translated_frags) > 1:
            trans_cols = f"{translated_frags[0]['columns']} ({len(translated_frags)} fragments)"
            trans_detail = " + ".join(f"P{t['page_index']+1}{t['rows_per_column']}" for t in translated_frags)
            trans_rows = f"{merged['rows_per_column'] if merged else ''} (Combined: {trans_detail})"
        elif translated_frags:
            trans_cols = translated_frags[0]["columns"]
            trans_rows = str(translated_frags[0]["rows_per_column"])
        else:
            trans_cols = ""
            trans_rows = ""

        row_values = [
            item.get("table_number"),
            display_status,
            item.get("reason"),
            item.get("sequence_page_number"),
            item.get("labeled_page_number"),
            eng_cols,
            eng_rows,
            trans_cols,
            trans_rows,
            "TRUE" if item.get("repeated_header") else "FALSE",
            "TRUE" if overflow else "FALSE",
            item.get("overflow_reasons_str", "")
        ]

        ws.append(row_values)

        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.border = thin_border

            if col_idx == 2:  # Status Column
                cell.font = status_font
                cell.alignment = align_center
            elif col_idx in [1, 4, 5, 6, 8, 10, 11]:
                cell.font = normal_font
                cell.alignment = align_center
            else:
                cell.font = normal_font
                cell.alignment = align_left

    # Auto column width adjustment
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    try:
        wb.save(REPORT_XLSX)
        print()
        print(f"Excel report (Colored):")
        print(REPORT_XLSX)
    except PermissionError:
        print(f"Warning: Could not write Excel report (file locked): {REPORT_XLSX}")


# ============================================================
# PRINT RESULT
# ============================================================

def print_results(
    results
):

    print()
    print("=" * 100)
    print("FINAL TABLE COMPARISON")
    print("=" * 100)

    passed = 0
    different = 0
    missing = 0
    extra = 0

    for result in results:

        number = result[
            "table_number"
        ]

        status = result[
            "status"
        ]

        if status == "PASS":
            passed += 1

        elif status == "DIFFERENT":
            different += 1

        elif status == "MISSING":
            missing += 1

        elif status == "EXTRA":
            extra += 1

        print()
        print(
            f"Table {number:04d}: "
            f"{status}"
        )

        print(
            f"    Reason: "
            f"{result.get('reason', '')}"
        )

        english = result.get(
            "english"
        )

        if english:

            print(
                f"    English:"
            )

            print(
                f"        "
                f"Columns = "
                f"{english['columns']}"
            )

            print(
                f"        "
                f"Rows/column = "
                f"{english['rows_per_column']}"
            )

        english_fragments = (
            result.get(
                "english_fragments",
                []
            )
        )

        if english_fragments:

            print(
                "    English fragments:"
            )

            for fragment in (
                english_fragments
            ):

                print(
                    f"        "
                    f"Page {fragment['page_index'] + 1}: "
                    f"{fragment['rows_per_column']}"
                )

        translated_fragments = (
            result.get(
                "translated_fragments",
                []
            )
        )

        if translated_fragments:

            print(
                "    Translated:"
            )

            for fragment in (
                translated_fragments
            ):

                print(
                    f"        "
                    f"Page {fragment['page_index'] + 1}: "
                    f"{fragment['rows_per_column']}"
                )

        if result.get(
            "merged"
        ):

            print(
                "    Reconstructed:"
            )

            print(
                f"        Columns = "
                f"{result['merged']['columns']}"
            )

            print(
                f"        Rows/column = "
                f"{result['merged']['rows_per_column']}"
            )

            print(
                f"        Repeated header = "
                f"{result.get('repeated_header', False)}"
            )

        # Check text overflow
        all_frags = []
        if result.get("english"):
            all_frags.append(result["english"])
        if result.get("english_fragments"):
            all_frags.extend(result["english_fragments"])
        if result.get("translated_fragments"):
            all_frags.extend(result["translated_fragments"])

        overflow_reasons = []
        for t in all_frags:
            if t and t.get("has_text_overflow"):
                overflow_reasons.extend(t.get("overflow_reasons", []))

        if overflow_reasons:
            unique_reasons = list(set(overflow_reasons))
            print(
                f"    [WARNING] Text Overflow Detected! ({'; '.join(unique_reasons)})"
            )

    print()
    print("=" * 100)
    print("SUMMARY")
    print("=" * 100)

    print(
        f"Compared   : "
        f"{len(results)}"
    )

    print(
        f"PASS       : "
        f"{passed}"
    )

    print(
        f"DIFFERENT  : "
        f"{different}"
    )

    print(
        f"MISSING    : "
        f"{missing}"
    )

    print(
        f"EXTRA      : "
        f"{extra}"
    )

    print("=" * 100)


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 100)
    print("OCR-FREE TABLE STRUCTURE COMPARISON")
    print("=" * 100)

    # ========================================================
    # ENGLISH
    # ========================================================

    print()
    print(
        "Analyzing English pages..."
    )

    english_tables = analyze_folder(
        ENGLISH_DIR
    )

    print()
    print(
        f"English physical tables: "
        f"{len(english_tables)}"
    )

    # ========================================================
    # TRANSLATED
    # ========================================================

    print()
    print(
        "Analyzing translated pages..."
    )

    translated_tables = analyze_folder(
        TRANSLATED_DIR
    )

    print()
    print(
        f"Translated physical tables: "
        f"{len(translated_tables)}"
    )

    # ========================================================
    # COMPARE
    # ========================================================

    results = compare_direction(
        english_tables,
        translated_tables
    )

    # ========================================================
    # PRINT
    # ========================================================

    print_results(
        results
    )

    # ========================================================
    # REPORT
    # ========================================================

    english_labels = load_pdf_page_labels(
        ENGLISH_PDF
    )

    translated_labels = load_pdf_page_labels(
        TRANSLATED_PDF
    )

    report = create_report(
        results,
        english_labels=english_labels,
        translated_labels=translated_labels
    )

    save_json(
        report
    )

    save_excel(
        report
    )

    print()
    print(
        f"JSON report:"
    )

    print(
        REPORT_JSON
    )

    print()
    print(
        f"Excel report (Colored):"
    )

    print(
        REPORT_XLSX
    )

    print()
    print(
        f"Debug images:"
    )

    print(
        DEBUG_DIR
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()