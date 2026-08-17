"""
main.py

Unified Entry Point for PDF Quality & Consistency Inspection.

Imports and coordinates:
  - FirstPage.py           : Checks Page 1 (Page Size, Version, Doc Length, Language Code, Title, Manual Type)
  - TOC.py                 : Checks Table of Contents Topic Numerics & Flags Missing Sections
  - LastPage.py            : Checks Last Page (Address, Disclaimer, Copyright, Language Code, Manual Type Code, Revision Date)
  - Barcode_QR_Check.py    : Multi-page Barcode & QR Code Count & Presence Validation (Non-visual)
  - crop_pdf_images.py     : Pure Graphic Element Extraction (excluding text & Barcode/QR)
  - Compare_cropped_images.py : Pure Visual Graphic Crops Comparison across pages

Generates:
  - PDF_Quality_Inspection_Report.xlsx : Unified multi-sheet Excel report (Overview, First Page, TOC, Last Page, Barcode & QR, Images)
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pymupdf

import FirstPage

try:
    import TOC
except ImportError:
    import Toc as TOC

import LastPage
import Barcode_QR_Check
import crop_pdf_images
import Compare_cropped_images


# ============================================================
# EXCEL REPORT GENERATORS
# ============================================================

def generate_unified_excel_report(
    fp_results,
    toc_results,
    lp_results,
    bc_qr_results,
    img_results_summary,
    img_crop_details_list,
    values_data,
    output_excel_path
):
    """
    Generate ONE single unified Excel report containing 6 worksheets:
      1. Overview    : Executive overview of all sub-check verdicts and Master Verdict
      2. First Page  : Detailed Page 1 checks (Title, Language Code, Manual Type, Doc Len, Version, Overall)
      3. TOC         : Detailed TOC topic numerics, missing topics, and Overall status
      4. Last Page   : Detailed Last Page checks (Address, Disclaimer, Copyright, Lang Code, Manual Type Code, Date, Overall)
      5. Barcode & QR: Dedicated Barcode & QR Code count and presence matching
      6. Images      : Detailed crop-by-crop visual graphic matching (Crop Name, Type, Eng Pg, Trans Pg, Movement, Similarity %, Status)
    """
    wb = openpyxl.Workbook()

    # Common Styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    pass_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="274E13")

    fail_fill = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
    fail_font = Font(name="Calibri", size=10, bold=True, color="783F04")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )

    # --------------------------------------------------------
    # TAB 1: OVERVIEW
    # --------------------------------------------------------
    ws_overview = wb.active
    ws_overview.title = "Overview"

    headers_overview = [
        "English Master PDF",
        "Translated PDF",
        "First Page",
        "TOC",
        "Last Page",
        "Images",
        "Barcode & QR",
        "Master Verdict",
    ]
    ws_overview.append(headers_overview)

    for i in range(len(fp_results)):
        fp_res = fp_results[i]
        toc_res = toc_results[i]
        lp_res = lp_results[i]
        bc_res = bc_qr_results[i]
        img_res = img_results_summary[i]

        fp_status = fp_res["overall_verdict"]
        toc_status = toc_res["status"]
        lp_status = lp_res["overall_verdict"]
        bc_status = bc_res["overall_verdict"]
        img_status = img_res["overall_status"]

        master_pass = (
            fp_status == "PASS" and
            toc_status == "PASS" and
            lp_status == "PASS" and
            bc_status == "PASS" and
            img_status == "PASS"
        )
        master_verdict = "PASS" if master_pass else "FAIL"

        row_data = [
            fp_res["english_pdf"],
            fp_res["translated_pdf"],
            fp_status,
            toc_status,
            lp_status,
            img_status,
            bc_status,
            master_verdict,
        ]
        ws_overview.append(row_data)

    # --------------------------------------------------------
    # TAB 2: FIRST PAGE (QR & Barcode columns removed)
    # --------------------------------------------------------
    ws_fp = wb.create_sheet(title="First Page")

    headers_fp = [
        "English Master PDF",
        "Translated PDF",
        "Title",
        "Language Code",
        "Manual Type",
        "Document Number Length",
        "Version",
        "Overall",
    ]
    ws_fp.append(headers_fp)

    for res in fp_results:
        row_data = [
            res["english_pdf"],
            res["translated_pdf"],
            res["title_status"],
            res["language_status"],
            res["manual_type_status"],
            res["doc_len_display"],
            res["version_display"],
            res["overall_verdict"],
        ]
        ws_fp.append(row_data)

    # --------------------------------------------------------
    # TAB 3: TOC
    # --------------------------------------------------------
    ws_toc = wb.create_sheet(title="TOC")

    headers_toc = [
        "English Master PDF",
        "Translated PDF",
        "TOC Topic Numerics",
        "Missing TOC Topic Numerics",
        "Overall",
    ]
    ws_toc.append(headers_toc)

    for i, res in enumerate(toc_results):
        vitem = values_data[i]
        row_data = [
            res["english_pdf"],
            res["translated_pdf"],
            vitem["toc_numerics"],
            res["diff_msg"],
            res["status"],
        ]
        ws_toc.append(row_data)

    # --------------------------------------------------------
    # TAB 4: LAST PAGE
    # --------------------------------------------------------
    ws_lp = wb.create_sheet(title="Last Page")

    headers_lp = [
        "English Master PDF",
        "Translated PDF",
        "Address",
        "DISCLAIMER",
        "Copyright",
        "Language Code",
        "Manual Type Code",
        "Revision Date",
        "Overall",
    ]
    ws_lp.append(headers_lp)

    for res in lp_results:
        row_data = [
            res["english_pdf"],
            res["translated_pdf"],
            res["address_status"],
            res["disclaimer_status"],
            res["copyright_status"],
            res["lang_code_status"],
            res["manual_type_code_status"],
            res["date_code_status"],
            res["overall_verdict"],
        ]
        ws_lp.append(row_data)

    # --------------------------------------------------------
    # TAB 5: BARCODE & QR (Dedicated Worksheet)
    # --------------------------------------------------------
    ws_bc = wb.create_sheet(title="Barcode & QR")

    headers_bc = [
        "English Master PDF",
        "Translated PDF",
        "Master Barcode Count",
        "Target Barcode Count",
        "Master Barcode Pages",
        "Target Barcode Pages",
        "Barcode Status",
        "Master QR Count",
        "Target QR Count",
        "Master QR Pages",
        "Target QR Pages",
        "QR Status",
        "Overall Status",
    ]
    ws_bc.append(headers_bc)

    for res in bc_qr_results:
        m_bc_pages = ", ".join(map(str, res.get("master_pages_barcode", []))) if res.get("master_pages_barcode") else "None"
        t_bc_pages = ", ".join(map(str, res.get("target_pages_barcode", []))) if res.get("target_pages_barcode") else "None"
        m_qr_pages = ", ".join(map(str, res.get("master_pages_qr", []))) if res.get("master_pages_qr") else "None"
        t_qr_pages = ", ".join(map(str, res.get("target_pages_qr", []))) if res.get("target_pages_qr") else "None"

        row_data = [
            res["english_pdf"],
            res["translated_pdf"],
            res["master_barcode_count"],
            res["target_barcode_count"],
            m_bc_pages,
            t_bc_pages,
            res["barcode_status"],
            res["master_qr_count"],
            res["target_qr_count"],
            m_qr_pages,
            t_qr_pages,
            res["qr_status"],
            res["overall_verdict"],
        ]
        ws_bc.append(row_data)

    # --------------------------------------------------------
    # TAB 6: IMAGES (Detailed Crop-by-Crop Worksheet)
    # --------------------------------------------------------
    ws_img = wb.create_sheet(title="Images")

    headers_img = [
        "English Master PDF",
        "Translated PDF",
        "Crop Name",
        "Image Type",
        "English Page",
        "Matched Target Page",
        "Page Movement",
        "Similarity (%)",
        "Match Status",
    ]
    ws_img.append(headers_img)

    for crop_row in img_crop_details_list:
        row_data = [
            crop_row["english_pdf"],
            crop_row["translated_pdf"],
            crop_row["crop_name"],
            crop_row["image_type"],
            crop_row["eng_page"],
            crop_row["trans_page"] if crop_row["trans_page"] != -1 else "N/A",
            crop_row["shift_info"],
            f"{crop_row['similarity']:.2f}%",
            crop_row["status"],
        ]
        ws_img.append(row_data)

    # Format all sheets with headers, borders, and PASS/Present/Equal/Matched fills
    for ws in wb.worksheets:
        # Style Header Row
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Style Data Rows
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="left" if (ws.title in ("TOC", "Images") and col_idx in (3, 4)) else "center",
                    vertical="center"
                )

                val = str(cell.value or "")
                if val.startswith("PASS") or val.startswith("Present") or val in ("Equal", "Matched", "MATCH (PASS)", "Same Page"):
                    cell.fill = pass_fill
                    cell.font = pass_font
                elif val.startswith("FAIL") or val in ("Not Present", "Not Equal", "Not Matched", "CHECK", "MISSING") or "Missing:" in val or "Extra:" in val:
                    cell.fill = fail_fill
                    cell.font = fail_font

        # Auto-fit Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            if ws.title == "TOC" and col[0].column in (3, 4):
                ws.column_dimensions[col_letter].width = 50
            elif ws.title == "Images" and col[0].column in (1, 2, 3):
                ws.column_dimensions[col_letter].width = 30
            else:
                ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)
    try:
        wb.save(output_excel_path)
        print(f"  [Unified Inspection Excel Report Saved]: {os.path.abspath(output_excel_path)}")
    except PermissionError:
        alt_path = output_excel_path.replace(".xlsx", "_new.xlsx")
        wb.save(alt_path)
        print(f"  [Unified Inspection Excel Report Saved]: {os.path.abspath(alt_path)} (locked file fallback)")


# ============================================================
# MASTER BATCH RUNNER
# ============================================================

def run_quality_inspection(source_pdf_path, translated_path_or_folder, output_dir):
    """
    Run Unified FirstPage, TOC, LastPage, Barcode_QR, and Images Inspection Report across translated PDFs.
    """
    if not os.path.exists(source_pdf_path):
        print(f"ERROR: Source English PDF not found: {source_pdf_path}")
        return

    print("=" * 80)
    print("UNIFIED PDF QUALITY & VISUAL INSPECTION ENGINE (MAIN)")
    print("=" * 80)
    print(f"Master English Source: {os.path.basename(source_pdf_path)}")
    print(f"Translated Target    : {translated_path_or_folder}")
    print(f"Output Directory     : {output_dir}")
    print("-" * 80)

    # 1. Discover translated PDFs
    translated_files = []
    if os.path.isfile(translated_path_or_folder):
        translated_files.append(translated_path_or_folder)
    elif os.path.isdir(translated_path_or_folder):
        for f in os.listdir(translated_path_or_folder):
            if f.lower().endswith(".pdf"):
                translated_files.append(os.path.join(translated_path_or_folder, f))
    else:
        print(f"ERROR: Invalid translated path or folder: {translated_path_or_folder}")
        return

    if not translated_files:
        print("ERROR: No translated PDF files found to inspect!")
        return

    # 2. Extract Master Source Models & Crop Graphic Elements
    print("Extracting Master Models for Source PDF...")
    source_fp_model = FirstPage.extract_page_model(source_pdf_path)
    source_toc_numerics = TOC.extract_toc_numerics(source_pdf_path)

    # Extract Master Last Page Footer Model
    doc_src = pymupdf.open(source_pdf_path)
    src_last_page_text = doc_src[-1].get_text("text")
    doc_src.close()
    src_last_lines = [l.strip() for l in src_last_page_text.splitlines() if l.strip()]
    src_footer_line = src_last_lines[-1] if src_last_lines else ""
    source_lp_model = LastPage.parse_footer_line(src_footer_line)

    # Crop pure graphic elements from Master English PDF for image comparison
    eng_crops_out_dir = os.path.join(output_dir, "Cropped_Images")
    pdf_name_no_ext = os.path.splitext(os.path.basename(source_pdf_path))[0]
    eng_crops_dir = os.path.join(eng_crops_out_dir, pdf_name_no_ext)

    print(f"Extracting Pure Graphic Crops from Source PDF...")
    crop_pdf_images.crop_pdf_elements(source_pdf_path, eng_crops_out_dir)

    print(f"  Source Language      : {source_fp_model['language']}")
    print(f"  Source Title         : {source_fp_model['title']} (max font: {source_fp_model['max_font_size']:.1f})")
    print(f"  Source MType (Page 1): {source_fp_model['manual_type']}")
    print(f"  Source Version       : {source_fp_model['version']}")
    print(f"  Source Doc Num       : {source_fp_model['document_number']} (len {source_fp_model['document_length']})")
    print(f"  Source Topics        : {len(source_toc_numerics)} sections")
    print(f"  Source Last Footer   : {source_lp_model['raw_line']}")
    print()

    # 3. Inspect each translated PDF across all modules
    print(f"Inspecting {len(translated_files)} Translated PDF(s)...")
    print()

    fp_results = []
    toc_results = []
    lp_results = []
    bc_qr_results = []
    img_results_summary = []
    img_crop_details_list = []
    values_data = []

    diff_crops_out_dir = os.path.join(output_dir, "Cropped_Comparison")

    for idx, tr_path in enumerate(translated_files, start=1):
        tr_filename = os.path.basename(tr_path)
        
        # 1. FirstPage check
        try:
            target_fp_model = FirstPage.extract_page_model(
                tr_path,
                ref_manual_bbox=source_fp_model.get("manual_type_bbox"),
                ref_manual_type=source_fp_model.get("manual_type")
            )
            fp_res = FirstPage.compare_page_models(source_fp_model, target_fp_model)
        except Exception as e:
            target_fp_model = {"language": "EN", "title": "Missing", "manual_type": "Others", "version": "N/A", "document_number": "N/A", "document_length": "N/A"}
            fp_res = {"english_pdf": os.path.basename(source_pdf_path), "translated_pdf": tr_filename, "language": "MISSING", "title_val": "MISSING", "title_status": "Not Equal", "manual_type_val": "Others", "manual_type_status": "Not Present", "page_size_val": "N/A", "page_size_status": "FAIL", "version_val": "N/A", "version_status": "FAIL", "version_display": "FAIL", "doc_len_val": "N/A", "doc_len_status": "FAIL", "doc_len_display": "FAIL", "language_status": "Not Present", "barcode_status": "Not Present", "qr_status": "Not Present", "overall_verdict": "FAIL"}
        fp_results.append(fp_res)

        # 2. TOC check
        try:
            target_toc_numerics = TOC.extract_toc_numerics(tr_path)
            toc_res = TOC.compare_toc_numerics(source_toc_numerics, target_toc_numerics)
            toc_res["english_pdf"] = os.path.basename(source_pdf_path)
            toc_res["translated_pdf"] = tr_filename
        except Exception as e:
            target_toc_numerics = []
            toc_res = {"english_pdf": os.path.basename(source_pdf_path), "translated_pdf": tr_filename, "diff_msg": str(e), "status": "FAIL"}
        toc_results.append(toc_res)

        # 3. LastPage check
        try:
            tr_lang = target_fp_model.get("language") or "EN"
            lp_res = LastPage.verify_last_page(
                tr_path,
                lang_code=tr_lang,
                ref_footer_model=source_lp_model,
                ref_manual_type=source_fp_model.get("manual_type")
            )
            lp_res["english_pdf"] = os.path.basename(source_pdf_path)
            lp_res["translated_pdf"] = tr_filename
        except Exception as e:
            lp_res = {"english_pdf": os.path.basename(source_pdf_path), "translated_pdf": tr_filename, "address_status": "FAIL", "disclaimer_status": "FAIL", "copyright_status": "FAIL", "lang_code_status": "FAIL", "manual_type_code_status": "FAIL", "date_code_status": "FAIL", "overall_verdict": "FAIL"}
        lp_results.append(lp_res)

        # 4. Barcode & QR Code Count Check
        try:
            bc_res = Barcode_QR_Check.compare_barcode_qr(source_pdf_path, tr_path)
        except Exception as e:
            bc_res = {"english_pdf": os.path.basename(source_pdf_path), "translated_pdf": tr_filename, "master_barcode_count": 0, "target_barcode_count": 0, "master_pages_barcode": [], "target_pages_barcode": [], "barcode_status": "FAIL", "master_qr_count": 0, "target_qr_count": 0, "master_pages_qr": [], "target_pages_qr": [], "qr_status": "FAIL", "overall_verdict": "FAIL"}
        bc_qr_results.append(bc_res)

        # 5. Pure Visual Graphic Images Comparison
        try:
            img_res = Compare_cropped_images.compare_english_crops_with_translated_pdf(
                eng_crop_dir=eng_crops_dir,
                trans_pdf_path=tr_path,
                output_dir=diff_crops_out_dir
            )
        except Exception as e:
            img_res = {"trans_name": tr_filename, "total_crops": 0, "matched_crops": 0, "match_pct": 0.0, "overall_status": "FAIL", "crop_details": []}
        img_results_summary.append(img_res)

        # Unpack per-crop details for Images tab
        for cd in img_res.get("crop_details", []):
            crop_name = cd["crop_file"]
            image_type = "table_image" if "table_image" in crop_name.lower() else "normal image"
            img_crop_details_list.append({
                "english_pdf": os.path.basename(source_pdf_path),
                "translated_pdf": tr_filename,
                "crop_name": crop_name,
                "image_type": image_type,
                "eng_page": cd["eng_page"],
                "trans_page": cd["trans_page"],
                "shift_info": cd["shift_info"],
                "similarity": cd["match_pct"],
                "status": cd["status"]
            })

        # Extracted Values Data
        values_data.append({
            "english_pdf": os.path.basename(source_pdf_path),
            "translated_pdf": tr_filename,
            "title_of_manual": target_fp_model.get("title", "Missing"),
            "manual_type": target_fp_model.get("manual_type", "Others"),
            "translated_manual_type": target_fp_model.get("translated_manual_text", "Missing") if target_fp_model.get("has_manual_type") else "Missing",
            "language_code": target_fp_model.get("language", "Missing"),
            "version": target_fp_model.get("version", "N/A"),
            "document_number": target_fp_model.get("document_number", "N/A"),
            "document_length": target_fp_model.get("document_length", "N/A"),
            "toc_numerics": ", ".join(target_toc_numerics) if target_toc_numerics else "None",
        })

        master_pass = (
            fp_res["overall_verdict"] == "PASS" and
            toc_res["status"] == "PASS" and
            lp_res["overall_verdict"] == "PASS" and
            bc_res["overall_verdict"] == "PASS" and
            img_res["overall_status"] == "PASS"
        )
        master_verdict = "PASS" if master_pass else "FAIL"

        print(
            f"  [{idx:02d}/{len(translated_files):02d}] {tr_filename[:30]:<30} | "
            f"Lang: {fp_res['language']:<4} | "
            f"FP: {fp_res['overall_verdict']:<4} | "
            f"TOC: {toc_res['status']:<4} | "
            f"LP: {lp_res['overall_verdict']:<4} | "
            f"BC/QR: {bc_res['overall_verdict']:<4} | "
            f"Img: {img_res['overall_status']:<4} | "
            f"Master: {master_verdict}"
        )

    print()
    print("-" * 80)

    # 4. Generate ONE single Excel report with 6 worksheets
    unified_report_path = os.path.join(output_dir, "PDF_Quality_Inspection_Report.xlsx")
    generate_unified_excel_report(
        fp_results,
        toc_results,
        lp_results,
        bc_qr_results,
        img_results_summary,
        img_crop_details_list,
        values_data,
        unified_report_path
    )

    print("=" * 80)
    print("UNIFIED QUALITY & VISUAL INSPECTION COMPLETE")
    print("=" * 80)


# ============================================================
# ENTRY POINT CONFIGURATION
# ============================================================

if __name__ == "__main__":
    if sys.stdout.encoding != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")

    default_eng_pdf = r"c:\Xylem Project\SpotCheck\Input\English\894387_5.0_en-US_2026-04_IOM.Start350.pdf"
    SOURCE_ENGLISH_PDF = default_eng_pdf

    default_tr_dir = r"c:\Xylem Project\SpotCheck\Input\Translated"
    TRANSLATED_TARGET = default_tr_dir

    default_out_dir = r"c:\Xylem Project\SpotCheck\Output"
    OUTPUT_DIRECTORY = default_out_dir

    run_quality_inspection(
        source_pdf_path=SOURCE_ENGLISH_PDF,
        translated_path_or_folder=TRANSLATED_TARGET,
        output_dir=OUTPUT_DIRECTORY
    )

